import docx
import json
import os
import sys

def extract_docx(file_path):
    doc = docx.Document(file_path)
    content = []
    
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            style_name = para.style.name if para.style else "Normal"
            content.append({"type": "paragraph", "style": style_name, "text": para.text})
            
    for i, table in enumerate(doc.tables):
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        content.append({"type": "table", "data": table_data})
        
    return content

def main():
    q_only = extract_docx("data/Questions_Only.docx")
    q_ans = extract_docx("data/Questions_and_Answers.docx")
    
    with open("data/extracted_questions_only.json", "w", encoding="utf-8") as f:
        json.dump(q_only, f, indent=4)
        
    with open("data/extracted_q_and_a.json", "w", encoding="utf-8") as f:
        json.dump(q_ans, f, indent=4)
        
if __name__ == "__main__":
    main()
